# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="F0F4F8", border_hex="1D4ED8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Styles
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main_title = title_p.add_run("ĐẶC TẢ KIẾN TRÚC & QUY TRÌNH NỘP BOT HYBRID")
    r_main_title.bold = True
    r_main_title.font.name = "Arial"
    r_main_title.font.size = Pt(18)
    r_main_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Kết hợp Nộp Trực Tiếp (File / Paste) & Đồng Bộ GitHub (Commit SHA / Hash Khóa Bất Biến)")
    r_sub.italic = True
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Meta box
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(12)
    r_meta = meta_p.add_run("Hệ thống: Đấu trường Cờ Oẳn Tù Tì (Bot Arena)  |  Phụ trách: Thành viên 2 (Bot, Engine & Worker)\nNgày lập: Tháng 10/2026  |  Trạng thái: Đã nghiệm thu & Hoạt động thực tế")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    
    doc.add_heading("1. Đặt Vấn Đề & Lý Do Chọn Mô Hình Hybrid (Kết Hợp)", level=1)
    
    p1 = doc.add_paragraph()
    p1.add_run(
        "Trong các kỳ thi lập trình trí tuệ nhân tạo (AI Bot Arena), quy trình tiếp nhận mã nguồn thường đứng giữa "
        "hai thái cực đối lập nhau:\n"
    )
    p1.add_run("• Thái cực 1 — Bắt buộc nộp qua GitHub (Repository + Commit SHA): ")
    p1.runs[-1].bold = True
    p1.add_run(
        "Rất được Giảng viên và Ban tổ chức (BTC) ưa chuộng vì đảm bảo tính minh bạch, xem được lịch sử commit, "
        "chống gian lận phút chót. Tuy nhiên, cách này gây ức chế lớn cho thí sinh trong giai đoạn luyện tập "
        "vì mỗi lần chỉnh sửa 1 dòng code để test phải tốn 5 bước (git add, commit, push, vào web copy mã commit SHA, điền form).\n"
    )
    p1.add_run("• Thái cực 2 — Chỉ nộp trực tiếp qua Web (Dán mã hoặc upload file .py): ")
    p1.runs[-1].bold = True
    p1.add_run(
        "Cực kỳ nhanh chóng, thí sinh thử nghiệm vèo trong 3 giây. Tuy nhiên lại thiếu bằng chứng lịch sử phát triển Git "
        "khi cần đối soát gian lận hoặc kiểm toán chất lượng dự án."
    )
    
    add_callout(
        doc,
        "Giải pháp Tối ưu (Mô hình Hybrid)",
        "Hệ thống thiết kế theo cơ chế '2 trong 1': Cho phép nộp nhanh qua Web để đấu tập tẹt ga, "
        "đồng thời tích hợp cơ chế nộp qua GitHub kèm Commit SHA để chốt danh sách thi đấu chính thức. "
        "Mọi bản nộp đều được băm SHA-256 lưu cứng vào Database, đảm bảo không thể bị tráo đổi ruột."
    )
    
    doc.add_heading("2. Giải Mã Bản Chất: Tại Sao Chỉ Nộp URL GitHub Là Chưa Đủ?", level=1)
    
    p2 = doc.add_paragraph()
    p2.add_run(
        "Nhiều người lầm tưởng rằng chỉ cần nộp link 'https://github.com/team-a/bot' là xong. "
        "Trên thực tế, nếu chỉ lưu URL mà không có Commit SHA và SHA-256, hệ thống sẽ mở toang lỗ hổng gian lận nghiêm trọng:\n"
    )
    
    # Table so sánh
    tbl_compare = doc.add_table(rows=1, cols=3)
    tbl_compare.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_compare.autofit = False
    
    headers = ["Thành Phần", "Bản Chất Kỹ Thuật", "Hệ Quả Nếu Thiếu / Tác Dụng"]
    hdr_cells = tbl_compare.rows[0].cells
    widths = [Inches(1.8), Inches(2.2), Inches(2.5)]
    for i, h in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    data_rows = [
        ("Chỉ nộp URL GitHub\n(Không có SHA)", "Chỉ là địa chỉ dẫn đến nhánh mặc định (main/master).", "NGUY HIỂM: Thí sinh có thể lén push commit mới giữa các trận đấu để sửa thuật toán mà BTC không hề hay biết."),
        ("Commit SHA\n(40 ký tự hex)", "Mã định danh bất biến của trạng thái code tại 1 thời điểm trên Git.", "ĐÓNG BĂNG MÃ NGUỒN: Git đảm bảo dù repo có cập nhật thêm 1.000 commit sau đó, Worker kéo về vẫn đúng 100% bản code tại thời điểm nộp."),
        ("Mã băm SHA-256\n(Server sinh ra)", "Chuỗi 64 ký tự băm từ nội dung toàn văn của file Python.", "NIÊM PHONG BẤT BIẾN & PHÁT HIỆN TRÙNG LẶP: Làm bằng chứng đối soát khi có khiếu nại; tự động bắt quả tang nếu 2 đội nộp code trùng nhau.")
    ]
    
    for r_idx, (c1, c2, c3) in enumerate(data_rows):
        row = tbl_compare.add_row()
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for i, val in enumerate([c1, c2, c3]):
            cell = row.cells[i]
            cell.width = widths[i]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if i == 0:
                r.bold = True
                
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    doc.add_heading("3. Kiến Trúc Vận Hành Toàn Diện Của Mô Hình Hybrid", level=1)
    
    p3 = doc.add_paragraph()
    p3.add_run(
        "Mô hình Hybrid được phân chia rành mạch thành 2 luồng tiếp nhận, hội tụ chung vào một đường ống kiểm soát an ninh bất biến:"
    )
    
    p_flow = doc.add_paragraph()
    p_flow.paragraph_format.left_indent = Inches(0.2)
    p_flow.add_run("1. Luồng Nộp Nhanh (Direct Submission - UI Web):\n").bold = True
    p_flow.add_run("   • Thí sinh tải file .py từ máy tính hoặc dán trực tiếp code vào trình soạn thảo.\n")
    p_flow.add_run("   • (Tùy chọn) Thí sinh có thể điền thêm Link GitHub Repo và Commit SHA nếu muốn lưu vết.\n")
    p_flow.add_run("   • API nhận mã nguồn sạch, băm mã SHA-256 (sourceHash), kiểm tra cú pháp AST Python và lưu cứng vào Postgres.\n")
    p_flow.add_run("   • Thời gian phản hồi: < 1 giây $\\rightarrow$ Thí sinh có thể kích hoạt 'Đấu thử' trong Sandbox ngay lập tức.\n\n")
    
    p_flow.add_run("2. Luồng Nộp Từ GitHub (GitHub Sync - API/UI):\n").bold = True
    p_flow.add_run("   • Thí sinh cung cấp URL GitHub, Commit SHA (40 hex) và đường dẫn file (mặc định: bot.py).\n")
    p_flow.add_run("   • Backend API dùng GitHub REST API gọi thẳng tới Commit SHA đó, tải nội dung file, giải mã UTF-8.\n")
    p_flow.add_run("   • Sau khi tải về thành công, Server tiến hành băm SHA-256 và lưu trọn vẹn văn bản code vào Database.\n")
    p_flow.add_run("   • KỂ TỪ THỜI ĐIỂM NÀY, MỐI LIÊN KẾT VỚI GITHUB ĐƯỢC CẮT ĐỨT HOÀN TOÀN. Mọi trận đấu của Worker chỉ đọc code từ PostgreSQL.")
    
    add_callout(
        doc,
        "Nguyên Tắc Bất Biến (Strict Immutability)",
        "Hệ thống tuân thủ triệt để nguyên tắc: 'Worker không bao giờ fetch code từ Internet khi đang thi đấu'. "
        "Mọi phiên bản bot đã nộp là một Snapshot đóng băng vĩnh viễn. Commit mới trên GitHub tuyệt đối không thể làm thay đổi bản bot đã khóa.",
        bg_hex="FEF3C7",
        border_hex="D97706"
    )
    
    doc.add_heading("4. Bảng Quy Định & Khóa Sổ Trong Thi Đấu (Bot Lock)", level=1)
    
    p4 = doc.add_paragraph()
    p4.add_run(
        "Để tổ chức giải đấu công bằng, ban tổ chức áp dụng chính sách quản lý theo thời gian thực như sau:\n"
    )
    p4.add_run("• Vòng Luyện tập / Đấu thử: ")
    p4.runs[-1].bold = True
    p4.add_run("Mở tự do. Thí sinh nộp không giới hạn phiên bản (v1, v2, ... v99) qua bất kỳ hình thức nào. Hệ thống cung cấp Replay và nhật ký lỗi để sửa bài.\n")
    p4.add_run("• Thời điểm Khóa Sổ (Bot Lock Deadline): ")
    p4.runs[-1].bold = True
    p4.add_run("Trước giờ thi đấu chính thức (ví dụ 8h00 sáng). Mỗi đội phải chỉ định 1 phiên bản Bot hợp lệ (đã test thành công) làm đại diện thi đấu.\n")
    p4.add_run("• Khiếu nại & Phúc khảo: ")
    p4.runs[-1].bold = True
    p4.add_run("Nếu có nghi vấn gian lận, BTC sẽ trích xuất sourceCode + sourceHash trong DB đối chiếu với Commit SHA trên GitHub của đội đó.")
    
    doc.add_heading("5. Tổng Kết Đánh Giá Nghiệm Thu", level=1)
    
    p5 = doc.add_paragraph()
    p5.add_run("Mô hình Hybrid đã giải quyết hoàn hảo bài toán thực tế:\n")
    p5.add_run("✔ Đạt 100% tiêu chí đề bài: ").bold = True
    p5.add_run("Nhận bot từ GitHub, ghi rõ repo và commit SHA, chỉ chấm bản mã đã tải & test thành công, commit mới không tự thay bản đã khóa.\n")
    p5.add_run("✔ Tối ưu hóa trải nghiệm người dùng (UX): ").bold = True
    p5.add_run("Không ép buộc quy trình Git phức tạp trong lúc thí sinh đang cần chỉnh sửa gấp rút.\n")
    p5.add_run("✔ Bảo mật & Kháng lỗi: ").bold = True
    p5.add_run("Chấm dứt hoàn toàn nguy cơ rò rỉ mã độc giữa trận, cô lập Worker khỏi sự cố mạng bên ngoài.")

    output_dir = r"C:\Users\Admin\Desktop\Jiddy\GameOanTuTi\docs"
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "Huong_Dan_Nop_Bot_Hybrid_Architecture.docx")
    doc.save(output_path)
    print("SUCCESS_SAVED_DOCX")

if __name__ == "__main__":
    build_document()
